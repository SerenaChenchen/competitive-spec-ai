from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font,Alignment
from openpyxl.utils import get_column_letter


def create_xlsx(results,docs,products,fields):
 wb=Workbook(); ws=wb.active;ws.title='Comparison Matrix'
 ws.append(['Specification']+products)
 lookup={(r.product_id,r.field):r for r in results}
 for field in fields:
  ws.append([field]+[(lookup[(p,field)].normalized_value if lookup[(p,field)].status in ('verified','manual_override') else ('⚠ '+lookup[(p,field)].raw_value if lookup[(p,field)].raw_value else '— '+lookup[(p,field)].status)) for p in products])
  for c,p in enumerate(products,2):
   r=lookup[(p,field)]
   if r.evidence:
    refs='; '.join(f'{e.filename} PDF p.{e.pdf_page} | {e.quote}' for e in r.evidence)
    from openpyxl.comments import Comment
    ws.cell(ws.max_row,c).comment=Comment(refs[:30000],'Evidence')
 ev=wb.create_sheet('Evidence & Sources');ev.append(['Product','Field','Raw Value','Status','File','PDF Page','Exact Quote','Quote Verified','Method'])
 quality=wb.create_sheet('Data Quality');quality.append(['Product','Field','Status','Note'])
 for r in results:
  if r.status!='verified':quality.append([r.product_id,r.field,r.status,r.note])
  for e in r.evidence:ev.append([r.product_id,r.field,r.raw_value,r.status,e.filename,e.pdf_page,e.quote,e.verified,e.method])
 summary=wb.create_sheet('Competitive Insights');summary.append(['Disclaimer','Automated comparative claims are withheld pending human review of extracted evidence.'])
 reg=wb.create_sheet('Source Registry');reg.append(['Filename','SHA-256','PDF Pages','Document ID'])
 for d in docs:reg.append([d['filename'],d['sha256'],d['page_count'],d['id']])
 log=wb.create_sheet('Change Log');log.append(['Product','Field','Original','Override','Note'])
 for r in results:
  if r.status=='manual_override':log.append([r.product_id,r.field,r.raw_value,r.normalized_value,r.note])
 for sheet in wb:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
  for c in sheet[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(bold=True,color='FFFFFF')
  for col in sheet.columns:
   ix=col[0].column
   width=max(13,min(60,max(len(str(cell.value or '')) for cell in col)+3))
   sheet.column_dimensions[get_column_letter(ix)].width=width
  for row in sheet:
   for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 out=BytesIO();wb.save(out);return out.getvalue()
